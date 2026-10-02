from io import BytesIO
from docx import Document
from docx.shared import Pt

def make_docx(text, title):
    doc = Document()
    p = doc.add_paragraph()
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(16)
    for line in text.splitlines():
        doc.add_paragraph(line)
    out = BytesIO()
    doc.save(out)
    return out.getvalue()

def make_pdf(text, title):
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm

    out = BytesIO()
    doc = SimpleDocTemplate(out, pagesize=A4, rightMargin=18*mm, leftMargin=18*mm,
                            topMargin=18*mm, bottomMargin=18*mm)
    styles = getSampleStyleSheet()
    story = [Paragraph(title, styles["Title"]), Spacer(1, 8)]
    for line in text.splitlines():
        if line.strip():
            story.append(Paragraph(line.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"),
                                   styles["BodyText"]))
            story.append(Spacer(1, 4))
    doc.build(story)
    return out.getvalue()
